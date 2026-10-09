#!/usr/bin/env python3
"""Generate standalone Hope Editions content from the Editions workbook.

This generator is intentionally independent of generate_directory_content.py.
It reads one Editions workbook and writes content.editions.js only after the
workbook passes the selected validation mode.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import os
import re
import sys
import tempfile
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable
from urllib.parse import urlparse

from openpyxl import load_workbook
from openpyxl.cell.cell import Cell


ALLOWED_LAYOUTS: dict[str, tuple[int, int]] = {
    "Two side by side": (2, 0),
    "Feature and one beside": (2, 1),
    "Three all small": (3, 0),
    "Feature and two beside": (3, 1),
    "Two on two": (4, 0),
    "Three then one": (4, 0),
    "Three above feature below": (4, 1),
    "Five all small": (5, 0),
    "Three above feature and one below": (5, 1),
}
GENERATED_FROM = "HA_Editions.xlsx"

REQUIRED_EDITION_FIELDS = (
    "Slug",
    "Edition Number",
    "Title",
    "Strapline",
    "Hero Paragraph 1",
    "Hero Paragraph 2",
    "Hero Paragraph 3",
    "Hero Sign-off",
    "Things Paragraph",
    "Painting URL",
    "Painting Alt",
    "Painting Artist",
    "Painting Title",
    "Painting Year",
    "Painting Medium",
    "Painting Collection",
    "Painting Rights",
    "Painting Line",
    "Pre-footer Heading",
    "Pre-footer Paragraph 1",
    "Pre-footer Paragraph 2",
    "Person Name",
    "Person Role",
)


def text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def normalise_label(value: Any) -> str:
    value = text(value).splitlines()[0] if text(value) else ""
    value = re.sub(r"\s*\([^)]*\)", "", value)
    value = re.sub(r"\s+", " ", value).strip().lower()
    return {"coming soon": "coming soon label"}.get(value, value)


def yes(value: Any) -> bool:
    return text(value).upper() == "YES"


def absolute_https(value: Any) -> bool:
    parsed = urlparse(text(value))
    return parsed.scheme == "https" and bool(parsed.netloc)


def normal_path(value: Any) -> str:
    path = text(value)
    if not path:
        return ""
    if not path.startswith("/"):
        return ""
    return path.rstrip("/") or "/"


def is_pale_yellow(cell: Cell) -> bool:
    """Return true only for the workbook's pale-yellow unresolved placeholder fill."""
    colour = cell.fill.fgColor
    rgb = text(colour.rgb).upper()
    return rgb.endswith("FFF2CC")


@dataclass(frozen=True)
class Issue:
    severity: str
    sheet: str
    cell: str
    field: str
    message: str


class Validation:
    def __init__(self) -> None:
        self.issues: list[Issue] = []

    def error(self, sheet: str, cell: str, field: str, message: str) -> None:
        self.issues.append(Issue("ERROR", sheet, cell, field, message))

    def warning(self, sheet: str, cell: str, field: str, message: str) -> None:
        self.issues.append(Issue("WARNING", sheet, cell, field, message))

    @property
    def errors(self) -> list[Issue]:
        return [issue for issue in self.issues if issue.severity == "ERROR"]


def headings(ws, required: Iterable[str]) -> tuple[int, dict[str, int]]:
    expected = {normalise_label(value) for value in required}
    for row in range(1, ws.max_row + 1):
        found: dict[str, int] = {}
        for column in range(1, ws.max_column + 1):
            label = normalise_label(ws.cell(row, column).value)
            if label:
                found[label] = column
        if expected.issubset(found):
            return row, found
    raise ValueError(f"Could not find a header row containing: {', '.join(required)}")


def cell_value(ws, row: int, mapping: dict[str, int], name: str) -> tuple[str, str]:
    column = mapping[normalise_label(name)]
    cell = ws.cell(row, column)
    return text(cell.value), cell.coordinate


def required_value(validation: Validation, ws, field_cells: dict[str, Cell], label: str) -> str:
    cell = field_cells.get(label)
    if cell is None:
        validation.error(ws.title, "—", label, "Required field label is missing from the EDITION section.")
        return ""
    value = text(cell.value)
    if not value:
        validation.error(ws.title, cell.coordinate, label, "Required value is blank.")
    elif is_pale_yellow(cell):
        validation.error(ws.title, cell.coordinate, label, "Required value still has the pale-yellow unresolved-placeholder fill.")
    return value


def parse_publication_switch(validation: Validation, ws) -> bool:
    expected = normalise_label("Publish prices to the website (YES/NO)")
    matches: list[Cell] = []
    for row in ws.iter_rows():
        for cell in row:
            if normalise_label(cell.value) == expected:
                matches.append(cell)
    if len(matches) != 1:
        validation.error(ws.title, "—", "Publish prices to the website (YES/NO)", "Expected exactly one publication-price setting label.")
        return False
    label_cell = matches[0]
    value_cell = ws.cell(label_cell.row, label_cell.column + 1)
    value = text(value_cell.value).upper()
    if value not in {"YES", "NO"}:
        validation.error(ws.title, value_cell.coordinate, "Publish prices to the website (YES/NO)", "Value must be exactly YES or NO.")
        return False
    return value == "YES"


def parse_roster(validation: Validation, ws) -> list[dict[str, Any]]:
    required = ("Active", "Slug", "Edition Number", "Title", "Page URL", "In site menu")
    try:
        header_row, mapping = headings(ws, required)
    except ValueError as exc:
        validation.error(ws.title, "—", "Editions Roster", str(exc))
        return []

    active: list[dict[str, Any]] = []
    seen_slugs: set[str] = set()
    seen_numbers: set[str] = set()
    seen_paths: set[str] = set()
    for row in range(header_row + 1, ws.max_row + 1):
        slug, slug_cell = cell_value(ws, row, mapping, "Slug")
        active_flag, active_cell = cell_value(ws, row, mapping, "Active")
        if not slug and not active_flag:
            continue
        if not yes(active_flag):
            continue
        number, number_cell = cell_value(ws, row, mapping, "Edition Number")
        title, title_cell = cell_value(ws, row, mapping, "Title")
        page_url, url_cell = cell_value(ws, row, mapping, "Page URL")
        in_menu, _ = cell_value(ws, row, mapping, "In site menu")
        for field, value, coordinate in (("Slug", slug, slug_cell), ("Edition Number", number, number_cell), ("Title", title, title_cell), ("Page URL", page_url, url_cell)):
            if not value:
                validation.error(ws.title, coordinate, field, "Required active roster value is blank.")
        path = normal_path(page_url)
        expected_path = "/editions/" + slug
        if path != expected_path:
            validation.error(ws.title, url_cell, "Page URL", f"Active Edition route must be exactly {expected_path}; found {page_url or 'blank'}.")
        for key, value, label, coordinate in (
            (seen_slugs, slug, "Slug", slug_cell),
            (seen_numbers, number, "Edition Number", number_cell),
            (seen_paths, path, "Page URL", url_cell),
        ):
            if value in key:
                validation.error(ws.title, coordinate, label, f"Duplicate active Edition {label.lower()}: {value}.")
            key.add(value)
        active.append({"slug": slug, "number": number, "title": title, "pageUrl": path, "inSiteMenu": yes(in_menu), "row": row})
    if not active:
        validation.error(ws.title, "—", "Editions Roster", "At least one active Edition is required.")
    return active


def edition_fields(ws, artists_header_row: int) -> dict[str, Cell]:
    fields: dict[str, Cell] = {}
    for row in range(1, artists_header_row):
        label = text(ws.cell(row, 1).value)
        if label:
            fields[re.sub(r"\s+", " ", label).strip()] = ws.cell(row, 2)
    return fields


def parse_artists(validation: Validation, ws, header_row: int, pieces_header_row: int) -> list[dict[str, Any]]:
    required = ("Active", "Artist Slug", "Artist Order", "Display Name", "Layout", "Why line")
    mapping = {normalise_label(ws.cell(header_row, col).value): col for col in range(1, ws.max_column + 1)}
    artists: list[dict[str, Any]] = []
    seen_orders: set[int] = set()
    seen_slugs: set[str] = set()
    for row in range(header_row + 1, pieces_header_row):
        slug, slug_cell = cell_value(ws, row, mapping, "Artist Slug")
        active_flag, active_cell = cell_value(ws, row, mapping, "Active")
        if not slug and not active_flag:
            continue
        if not yes(active_flag):
            continue
        order_value, order_cell = cell_value(ws, row, mapping, "Artist Order")
        display_name, name_cell = cell_value(ws, row, mapping, "Display Name")
        layout, layout_cell = cell_value(ws, row, mapping, "Layout")
        why_line, why_cell = cell_value(ws, row, mapping, "Why line")
        for field, value, coordinate in (("Artist Slug", slug, slug_cell), ("Artist Order", order_value, order_cell), ("Display Name", display_name, name_cell), ("Layout", layout, layout_cell), ("Why line", why_line, why_cell)):
            if not value:
                validation.error(ws.title, coordinate, field, "Required active artist value is blank.")
        try:
            order = int(order_value)
            if str(order) != order_value and not order_value.endswith(".0"):
                raise ValueError
            if order <= 0:
                raise ValueError
        except ValueError:
            validation.error(ws.title, order_cell, "Artist Order", "Artist Order must be a positive whole number.")
            order = 0
        if order in seen_orders:
            validation.error(ws.title, order_cell, "Artist Order", f"Duplicate active Artist Order: {order}.")
        seen_orders.add(order)
        if slug in seen_slugs:
            validation.error(ws.title, slug_cell, "Artist Slug", f"Duplicate active Artist Slug: {slug}.")
        seen_slugs.add(slug)
        if layout not in ALLOWED_LAYOUTS:
            validation.error(ws.title, layout_cell, "Layout", f"Layout must be one of: {', '.join(ALLOWED_LAYOUTS)}.")
        artists.append({"slug": slug, "order": order, "displayName": display_name, "layout": layout, "whyLine": why_line, "row": row})
    return sorted(artists, key=lambda item: (item["order"], item["slug"]))


def parse_pieces(validation: Validation, ws, header_row: int, active_artists: dict[str, dict[str, Any]], publish_prices: bool) -> dict[str, list[dict[str, Any]]]:
    required = ("Active", "Artist Slug", "PieceOrder", "Piece Title", "World", "Status", "Medium", "Price", "Listing URL", "Featured", "Meaning", "Good For 1", "Good For 2", "Good For 3", "Image URL", "Image Alt", "Coming Soon Label")
    mapping = {normalise_label(ws.cell(header_row, col).value): col for col in range(1, ws.max_column + 1)}
    if not all(normalise_label(label) in mapping for label in required):
        validation.error(ws.title, "—", "PIECES", "PIECES headers are incomplete or renamed.")
        return {artist_slug: [] for artist_slug in active_artists}

    by_artist: dict[str, list[dict[str, Any]]] = {artist_slug: [] for artist_slug in active_artists}
    seen_orders: dict[str, set[int]] = {artist_slug: set() for artist_slug in active_artists}
    for row in range(header_row + 1, ws.max_row + 1):
        artist_slug, artist_cell = cell_value(ws, row, mapping, "Artist Slug")
        active_flag, active_cell = cell_value(ws, row, mapping, "Active")
        status, status_cell = cell_value(ws, row, mapping, "Status")
        if not artist_slug and not active_flag and not status:
            continue
        if artist_slug not in active_artists or not yes(active_flag) or status.lower() != "available":
            continue
        order_value, order_cell = cell_value(ws, row, mapping, "PieceOrder")
        title, title_cell = cell_value(ws, row, mapping, "Piece Title")
        world, world_cell = cell_value(ws, row, mapping, "World")
        medium, _ = cell_value(ws, row, mapping, "Medium")
        price, price_cell = cell_value(ws, row, mapping, "Price")
        listing_url, listing_cell = cell_value(ws, row, mapping, "Listing URL")
        featured_value, _ = cell_value(ws, row, mapping, "Featured")
        meaning, _ = cell_value(ws, row, mapping, "Meaning")
        image_url, image_url_cell = cell_value(ws, row, mapping, "Image URL")
        image_alt, image_alt_cell = cell_value(ws, row, mapping, "Image Alt")
        coming_label, label_cell = cell_value(ws, row, mapping, "Coming Soon Label")
        good_for = [cell_value(ws, row, mapping, label)[0] for label in ("Good For 1", "Good For 2", "Good For 3")]

        for field, value, coordinate in (("PieceOrder", order_value, order_cell), ("Piece Title", title, title_cell), ("World", world, world_cell), ("Listing URL", listing_url, listing_cell), ("Image URL", image_url, image_url_cell), ("Image Alt", image_alt, image_alt_cell)):
            if not value:
                validation.error(ws.title, coordinate, field, "Required active available piece value is blank.")
            elif is_pale_yellow(ws[coordinate]):
                validation.error(ws.title, coordinate, field, "Required active available piece value still has the pale-yellow unresolved-placeholder fill.")
        try:
            order = int(order_value)
            if str(order) != order_value and not order_value.endswith(".0"):
                raise ValueError
            if order <= 0:
                raise ValueError
        except ValueError:
            validation.error(ws.title, order_cell, "PieceOrder", "PieceOrder must be a positive whole number.")
            order = 0
        if order in seen_orders[artist_slug]:
            validation.error(ws.title, order_cell, "PieceOrder", f"Duplicate active available PieceOrder for {artist_slug}: {order}.")
        seen_orders[artist_slug].add(order)
        if world not in {"To Keep", "To Make"}:
            validation.error(ws.title, world_cell, "World", "World must be exactly To Keep or To Make.")
        if listing_url and not absolute_https(listing_url):
            validation.error(ws.title, listing_cell, "Listing URL", "Listing URL must be an absolute https:// URL.")
        if image_url and not absolute_https(image_url):
            validation.error(ws.title, image_url_cell, "Image URL", "Image URL must be an absolute https:// URL.")
        if coming_label and coming_label != "Making more":
            validation.error(ws.title, label_cell, "Coming Soon Label", "Published cards may use only the controlled label Making more.")
        if publish_prices:
            try:
                numeric_price = float(price)
                if numeric_price <= 0:
                    raise ValueError
            except ValueError:
                validation.error(ws.title, price_cell, "Price (£) number only", "A positive numeric price is required when publication prices are YES.")
                numeric_price = 0.0
        else:
            numeric_price = None
        piece: dict[str, Any] = {
            "order": order,
            "title": title,
            "world": world,
            "medium": medium,
            "listingUrl": listing_url,
            "featured": yes(featured_value),
            "meaning": meaning,
            "goodFor": [item for item in good_for if item],
            "imageUrl": image_url,
            "imageAlt": image_alt,
            "comingSoonLabel": coming_label,
        }
        if publish_prices:
            piece["price"] = int(numeric_price) if numeric_price.is_integer() else numeric_price
        by_artist[artist_slug].append(piece)

    for artist_slug, pieces in by_artist.items():
        pieces.sort(key=lambda item: (item["order"], item["title"]))
        artist = active_artists[artist_slug]
        layout = artist["layout"]
        expected = ALLOWED_LAYOUTS.get(layout)
        if expected:
            expected_count, expected_featured = expected
            if len(pieces) != expected_count:
                validation.error(ws.title, "—", f"Pieces for {artist_slug}", f"Layout {layout!r} needs {expected_count} active available pieces; found {len(pieces)}.")
            actual_featured = sum(1 for piece in pieces if piece["featured"])
            if actual_featured != expected_featured:
                validation.error(ws.title, "—", f"Featured pieces for {artist_slug}", f"Layout {layout!r} needs {expected_featured} featured piece(s); found {actual_featured}.")
    return by_artist


def validate_pullouts(validation: Validation, ws, field_cells: dict[str, Cell]) -> None:
    count = 0
    for label in ("Hero Paragraph 1", "Hero Paragraph 2"):
        cell = field_cells.get(label)
        value = text(cell.value) if cell else ""
        if value.count("*") % 2:
            validation.error(ws.title, cell.coordinate if cell else "—", label, "Asterisk pull-out markers are unmatched.")
            continue
        matches = re.findall(r"\*([^*]+)\*", value)
        if any(not phrase.strip() for phrase in matches):
            validation.error(ws.title, cell.coordinate if cell else "—", label, "Asterisk pull-out phrases must not be empty.")
        count += len(matches)
    if count > 2:
        validation.error(ws.title, "—", "Hero Paragraphs 1–2", f"At most two pull-out phrases are allowed; found {count}.")


def parse_edition(validation: Validation, workbook, roster: dict[str, Any], publish_prices: bool) -> dict[str, Any] | None:
    slug = roster["slug"]
    if slug not in workbook.sheetnames:
        validation.error("Editions Roster", "—", "Slug", f"Active Edition tab {slug!r} is missing.")
        return None
    ws = workbook[slug]
    try:
        artists_header_row, _ = headings(ws, ("Active", "Artist Slug", "Artist Order", "Display Name", "Layout", "Why line"))
        pieces_header_row, _ = headings(ws, ("Active", "Artist Slug", "PieceOrder", "Piece Title", "World", "Status", "Listing URL", "Image URL", "Image Alt"))
    except ValueError as exc:
        validation.error(ws.title, "—", "Edition section", str(exc))
        return None
    fields = edition_fields(ws, artists_header_row)
    values = {field: required_value(validation, ws, fields, field) for field in REQUIRED_EDITION_FIELDS}

    meta_cell = fields.get("Meta Description")
    meta_description = text(meta_cell.value) if meta_cell else ""
    if not meta_description:
        validation.warning(ws.title, meta_cell.coordinate if meta_cell else "—", "Meta Description", "Meta Description is blank. It is optional and is not rendered; add it in Squarespace SEO settings when ready.")

    photo_url_cell = fields.get("Photo URL")
    photo_alt_cell = fields.get("Photo Alt")
    photo_url = text(photo_url_cell.value) if photo_url_cell else ""
    photo_alt = text(photo_alt_cell.value) if photo_alt_cell else ""
    if bool(photo_url) != bool(photo_alt):
        missing_cell = photo_alt_cell if photo_url else photo_url_cell
        validation.error(ws.title, missing_cell.coordinate if missing_cell else "—", "Pre-footer photo", "Photo URL and Photo Alt are an optional pair: fill both or leave both blank.")
    elif photo_url and not absolute_https(photo_url):
        validation.error(ws.title, photo_url_cell.coordinate, "Photo URL", "Photo URL must be an absolute https:// URL.")
    elif not photo_url:
        validation.warning(ws.title, photo_url_cell.coordinate if photo_url_cell else "—", "Pre-footer photo", "Photo URL and Photo Alt are both blank; the pre-footer will render without an image.")

    partner_cells = {label: fields.get(label) for label in ("Partner Name", "Partner Logo URL", "Partner Logo Alt")}
    partner_values = {label: text(cell.value) if cell else "" for label, cell in partner_cells.items()}
    partner_filled = [label for label, value in partner_values.items() if value]
    if partner_filled and len(partner_filled) != 3:
        first = partner_cells[partner_filled[0]]
        validation.error(ws.title, first.coordinate if first else "—", "Partner", "Partner Name, Partner Logo URL, and Partner Logo Alt must be all filled or all blank.")
    if partner_values["Partner Logo URL"] and not absolute_https(partner_values["Partner Logo URL"]):
        cell = partner_cells["Partner Logo URL"]
        validation.error(ws.title, cell.coordinate if cell else "—", "Partner Logo URL", "Partner Logo URL must be an absolute https:// URL.")

    if values["Slug"] != slug:
        cell = fields.get("Slug")
        validation.error(ws.title, cell.coordinate if cell else "—", "Slug", f"Edition tab Slug must match roster slug {slug!r}.")
    if values["Edition Number"] != roster["number"]:
        cell = fields.get("Edition Number")
        validation.error(ws.title, cell.coordinate if cell else "—", "Edition Number", "Edition tab number must match its roster row.")
    if values["Title"] != roster["title"]:
        cell = fields.get("Title")
        validation.error(ws.title, cell.coordinate if cell else "—", "Title", "Edition tab title must match its roster row.")
    validate_pullouts(validation, ws, fields)

    artists = parse_artists(validation, ws, artists_header_row, pieces_header_row)
    active_artists = {artist["slug"]: artist for artist in artists}
    pieces_by_artist = parse_pieces(validation, ws, pieces_header_row, active_artists, publish_prices)
    rendered_artists: list[dict[str, Any]] = []
    all_pieces: list[dict[str, Any]] = []
    for artist in artists:
        row_pieces = pieces_by_artist.get(artist["slug"], [])
        rendered = {key: value for key, value in artist.items() if key != "row"}
        rendered["pieces"] = row_pieces
        rendered_artists.append(rendered)
        all_pieces.extend(row_pieces)

    stats = {
        "artists": len(rendered_artists),
        "toKeep": sum(1 for piece in all_pieces if piece["world"] == "To Keep"),
        "toMake": sum(1 for piece in all_pieces if piece["world"] == "To Make"),
    }
    partner = None
    if len(partner_filled) == 3:
        partner = {"name": partner_values["Partner Name"], "logoUrl": partner_values["Partner Logo URL"], "logoAlt": partner_values["Partner Logo Alt"]}
    photo = {"url": photo_url, "alt": photo_alt} if photo_url and photo_alt else None
    return {
        "slug": slug,
        "editionNumber": values["Edition Number"],
        "title": values["Title"],
        "pageUrl": roster["pageUrl"],
        "inSiteMenu": roster["inSiteMenu"],
        "metaDescription": meta_description,
        "hero": {
            "strapline": values["Strapline"],
            "paragraph1": values["Hero Paragraph 1"],
            "paragraph2": values["Hero Paragraph 2"],
            "paragraph3": values["Hero Paragraph 3"],
            "signoff": values["Hero Sign-off"],
            "thingsParagraph": values["Things Paragraph"],
        },
        "painting": {
            "url": values["Painting URL"],
            "alt": values["Painting Alt"],
            "artist": values["Painting Artist"],
            "title": values["Painting Title"],
            "year": values["Painting Year"],
            "medium": values["Painting Medium"],
            "collection": values["Painting Collection"],
            "rights": values["Painting Rights"],
            "line": values["Painting Line"],
        },
        "partner": partner,
        "preFooter": {
            "heading": values["Pre-footer Heading"],
            "paragraph1": values["Pre-footer Paragraph 1"],
            "paragraph2": values["Pre-footer Paragraph 2"],
            "photo": photo,
            "personName": values["Person Name"],
            "personRole": values["Person Role"],
        },
        "stats": stats,
        "artists": rendered_artists,
    }


def build_payload(workbook_path: Path, validation: Validation) -> dict[str, Any]:
    workbook = load_workbook(workbook_path, data_only=True)
    if "README" not in workbook.sheetnames:
        validation.error("Workbook", "—", "README", "README sheet is missing.")
        return {"generatedFrom": GENERATED_FROM, "publishPrices": False, "editions": {}}
    if "Editions Roster" not in workbook.sheetnames:
        validation.error("Workbook", "—", "Editions Roster", "Editions Roster sheet is missing.")
        return {"generatedFrom": GENERATED_FROM, "publishPrices": False, "editions": {}}
    publish_prices = parse_publication_switch(validation, workbook["README"])
    roster = parse_roster(validation, workbook["Editions Roster"])
    editions: dict[str, Any] = {}
    for entry in roster:
        edition = parse_edition(validation, workbook, entry, publish_prices)
        if edition:
            editions[entry["slug"]] = edition
    return {"generatedFrom": GENERATED_FROM, "publishPrices": publish_prices, "editions": editions}


def js_source(payload: dict[str, Any]) -> str:
    body = json.dumps(payload, ensure_ascii=False, indent=2, separators=(",", ": "))
    return (
        "// Auto-generated by generate_editions_content.py from HA_Editions.xlsx.\n"
        "// DO NOT EDIT MANUALLY — amend the workbook and regenerate.\n"
        "window.HA_EDITIONS_CONTENT = " + body + ";\n"
    )


def report_markdown(workbook_path: Path, payload: dict[str, Any], validation: Validation, mode: str) -> str:
    errors = validation.errors
    warnings = [issue for issue in validation.issues if issue.severity == "WARNING"]
    digest = hashlib.sha256(workbook_path.read_bytes()).hexdigest()
    lines = [
        "# Hope Editions generation report",
        "",
        f"**Generated at:** {datetime.now(timezone.utc).strftime('%Y-%m-%dT%H:%M:%SZ')}",
        f"**Workbook:** `{workbook_path.name}`",
        f"**Workbook SHA-256:** `{digest}`",
        f"**Mode:** `{mode}`",
        f"**Result:** {'PASS' if not errors else 'FAIL'}",
        "",
        "## Publication settings",
        "",
        f"- Publish prices to the website: `{'YES' if payload.get('publishPrices') else 'NO'}`",
        "- Public price fields emitted: " + ("yes" if payload.get("publishPrices") else "no"),
        "",
        "## Generated Edition summary",
        "",
        "| Slug | Route | Artists | To Keep pieces | To Make pieces | Pre-footer photo |",
        "| --- | --- | ---: | ---: | ---: | --- |",
    ]
    for edition in payload.get("editions", {}).values():
        stats = edition["stats"]
        photo = "present" if edition["preFooter"]["photo"] else "omitted (blank optional pair)"
        lines.append(f"| `{edition['slug']}` | `{edition['pageUrl']}` | {stats['artists']} | {stats['toKeep']} | {stats['toMake']} | {photo} |")
    if not payload.get("editions"):
        lines.append("| — | — | 0 | 0 | 0 | — |")
    for heading, issues, empty in (
        ("Errors", errors, "No errors."),
        ("Warnings", warnings, "No warnings."),
    ):
        lines.extend(["", f"## {heading}", "", "| Sheet | Cell | Field | Message |", "| --- | --- | --- | --- |"])
        if issues:
            lines.extend(f"| `{issue.sheet}` | `{issue.cell}` | {issue.field} | {issue.message} |" for issue in issues)
        else:
            lines.append(f"| — | — | — | {empty} |")
    lines.extend([
        "",
        "## Output rule",
        "",
        "Strict generation writes `content.editions.js` only when there are no errors. Blank Meta Description is a warning only. Blank Photo URL plus blank Photo Alt is an allowed pair and renders no photo box; exactly one populated photo field is a strict error.",
        "",
    ])
    return "\n".join(lines)


def atomic_write(path: Path, content: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, temporary = tempfile.mkstemp(prefix=f".{path.name}.", suffix=".tmp", dir=path.parent)
    try:
        with os.fdopen(fd, "w", encoding="utf-8", newline="\n") as handle:
            handle.write(content)
        os.replace(temporary, path)
    except Exception:
        try:
            os.unlink(temporary)
        except OSError:
            pass
        raise


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", type=Path, required=True, help="Path to HA_Editions.xlsx.")
    parser.add_argument("--output", type=Path, required=True, help="Output content.editions.js path.")
    parser.add_argument("--report", type=Path, required=True, help="Markdown generation report path.")
    parser.add_argument("--mode", choices=("strict", "permissive"), default="strict")
    args = parser.parse_args()
    if not args.workbook.is_file():
        print(f"Workbook does not exist: {args.workbook}", file=sys.stderr)
        return 2
    validation = Validation()
    payload = build_payload(args.workbook, validation)
    report = report_markdown(args.workbook, payload, validation, args.mode)
    atomic_write(args.report, report)
    if validation.errors and args.mode == "strict":
        return 1
    if validation.errors:
        return 1
    atomic_write(args.output, js_source(payload))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
